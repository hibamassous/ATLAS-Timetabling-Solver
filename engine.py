# -*- coding: utf-8 -*-
"""
Moteur métier pour l'application "Emploi du temps & Affectation des salles".

Ce module adapte la logique de l'algorithme génétique (allocator.py) au
format attendu par app.py : des cours représentés par des dictionnaires
(num_edt, filiere, jour, debut, fin, duree, module, prof, effectif, salle),
un moteur pilotable en arrière-plan avec callbacks de progression, et la
génération des classeurs Excel d'export.
"""

import random
import statistics
import threading
import time
import uuid
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import allocator as core

# ============================================
# CONSTANTES PARTAGÉES AVEC LE FRONTEND
# ============================================
JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi"]
H_MIN, H_MAX = 8, 19
HEURES = list(range(H_MIN, H_MAX))

# Valeurs par defaut au demarrage (catalogue par defaut) — la source de
# verite courante, modifiable par l'administration, est get_salles().
SALLES_AMPHI = core.SALLES_AMPHI
SALLES_MOYENNE = core.SALLES_MOYENNE
SALLES_PETITE = core.SALLES_PETITE
SALLES_ALL = core.SALLES_ALL
CAPACITES = core.CAPACITES

UNIVERSITE = "Université Mohammed V, Rabat"


# ============================================
# CATALOGUE DES SALLES (dynamique, modifiable par l'admin)
# ============================================
def get_salles():
    """Retourne (amphis, moyennes, petites, toutes, capacites) a partir du
    catalogue courant (rooms_config.json), modifiable via l'administration."""
    rooms = core.load_or_init_rooms()
    amphi = [r.name for r in rooms if r.type == "Amphi"]
    moyenne = [r.name for r in rooms if "moyenne" in r.type.lower() and r.type != "Amphi"]
    petite = [r.name for r in rooms if r.name not in amphi and r.name not in moyenne]
    capacites = {r.name: r.capacity for r in rooms}
    return amphi, moyenne, petite, amphi + moyenne + petite, capacites


def rooms_payload():
    return core.list_rooms_payload()


def upsert_room(payload):
    return core.upsert_room(payload)


def remove_room(name):
    return core.remove_room(name)


# ============================================
# CHARGEMENT DU FICHIER EXCEL DE COURS
# ============================================
import re
import unicodedata

# Alias reconnus pour chaque role de colonne (compares apres normalisation :
# minuscules + accents retires). L'ordre des colonnes dans le fichier n'a
# aucune importance, seul l'intitule de l'en-tete compte.
_HEADER_ALIASES = {
    "num_edt": ["edt", "groupe", "group"],
    "filiere": ["filiere", "specialite", "parcours"],
    "module": ["module", "cours", "matiere", "subject"],
    "prof": ["professeur", "prof", "enseignant", "teacher"],
    "jour": ["jour", "day"],
    "horaire": ["horaire", "heure", "creneau", "time"],
    "effectif": ["effectif", "students", "nb etudiants", "nombre etudiants", "nb_etudiants"],
}

_TIME_RE = re.compile(r"(\d{1,2})\s*[:hH]\s*(\d{2})?")


def _norm(s):
    """Minuscule + accents retires, pour une comparaison robuste des en-tetes."""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.strip().lower()


def _find_header_row_and_map(df):
    """Cherche, parmi les 10 premieres lignes, celle qui ressemble le plus a
    une ligne d'en-tetes, et construit la correspondance colonne -> role."""
    best_row, best_score, best_map = None, 0, {}
    for idx in range(min(10, len(df))):
        row = df.iloc[idx]
        col_map = {}
        for col_idx, cell in enumerate(row):
            if not isinstance(cell, str) or not cell.strip():
                continue
            norm = _norm(cell)
            for role, aliases in _HEADER_ALIASES.items():
                if role in col_map.values():
                    continue
                if any(alias in norm for alias in aliases):
                    col_map[col_idx] = role
                    break
        score = len(set(col_map.values()))
        if score > best_score:
            best_row, best_score, best_map = idx, score, col_map
    return best_row, best_map, best_score


def _parse_horaire(raw):
    """Extrait (debut, fin) en heures pleines depuis un texte d'horaire,
    quel que soit le separateur utilise (-, en tiret –, em tiret —, "a", "/")."""
    matches = _TIME_RE.findall(str(raw))
    if len(matches) < 2:
        return 0, 0
    try:
        return int(matches[0][0]), int(matches[1][0])
    except (TypeError, ValueError):
        return 0, 0


def _detect_priority(module_name):
    u = str(module_name).upper()
    if "CM" in u:
        return 10
    if "TD" in u:
        return 8
    if "TP" in u:
        return 7
    if "EXAM" in u:
        return 9
    return 5


def charger_cours(path):
    """Lit un fichier Excel de cours et retourne (liste_de_cours_dict,
    liste_avertissements). Les colonnes sont detectees par leur en-tete
    (N° EDT / Filiere / Module / Professeur / Jour / Horaire / Effectif),
    peu importe leur ordre dans le fichier."""
    df = pd.read_excel(path, header=None)
    header_row, col_map, score = _find_header_row_and_map(df)

    warnings = []
    if header_row is None or score < 3:
        raise ValueError(
            "Unable to recognize the file's columns. Expected at "
            "minimum: EDT No. (or Group), Day, Time, Module (or Course), "
            "Class size. Check the header row.")

    role_to_col = {role: idx for idx, role in col_map.items()}
    missing = [r for r in ("num_edt", "jour", "horaire", "module") if r not in role_to_col]
    if missing:
        raise ValueError(
            "Missing columns in the file: " + ", ".join(missing) + ".")

    cours = []
    for idx in range(header_row + 1, len(df)):
        row = df.iloc[idx]

        def get(role, default=""):
            col = role_to_col.get(role)
            if col is None or col >= len(row) or pd.isna(row[col]):
                return default
            return row[col]

        num_edt = str(get("num_edt", "")).strip()
        if not num_edt or _norm(num_edt) in ("nan", "edt", "n edt", "n° edt", "groupe"):
            continue

        module = str(get("module", "")).strip()
        prof = str(get("prof", "")).strip()
        jour = str(get("jour", "")).strip()
        horaire_raw = get("horaire", "")
        filiere = str(get("filiere", "")).strip() or num_edt
        effectif_raw = get("effectif", 0)

        if not module or not jour or not horaire_raw:
            continue

        # Si aucune colonne "Professeur" distincte n'a ete trouvee, on tente
        # de la retrouver dans le libelle du module ("Module - Professeur").
        if not prof and " - " in module:
            module, prof = [p.strip() for p in module.split(" - ", 1)]
        prof = prof or "TBD"

        debut, fin = _parse_horaire(horaire_raw)
        if debut == 0 and fin == 0:
            warnings.append(
                f"Unreadable time slot for {num_edt} / {module} (\"{horaire_raw}\").")

        if jour not in JOURS:
            warnings.append(
                f"Unrecognized day \"{jour}\" for {num_edt} / {module} "
                "— this course will not appear in the day-by-day grids.")

        try:
            effectif = int(float(effectif_raw))
        except (TypeError, ValueError):
            effectif = 0

        cours.append({
            "id": uuid.uuid4().hex[:10],
            "num_edt": num_edt,
            "filiere": filiere,
            "jour": jour,
            "debut": debut,
            "fin": fin,
            "duree": max(fin - debut, 0),
            "module": module,
            "prof": prof,
            "effectif": effectif,
            "salle": None,
            "_priority": _detect_priority(module),
        })

    if not cours:
        warnings.append(
            "No courses recognized in the file. Check the columns "
            "(EDT No., Program, Module, Professor, Day, Time, Class size).")
    return cours, warnings


# ============================================
# GESTION DES COURS (ajout / modification / suppression manuelle)
# ============================================
def add_course(cours, payload):
    """Ajoute un cours (nouveau ou en complement d'un fichier importe) a la
    liste `cours` (en place). Retourne le cours cree."""
    num_edt = str(payload.get("num_edt", "")).strip()
    module = str(payload.get("module", "")).strip()
    prof = str(payload.get("prof", "")).strip() or "TBD"
    jour = str(payload.get("jour", "")).strip()
    filiere = str(payload.get("filiere", "")).strip() or num_edt

    if not num_edt:
        raise ValueError("The EDT No. / group is required.")
    if not module:
        raise ValueError("The module name is required.")
    if jour not in JOURS:
        raise ValueError("Invalid day.")

    try:
        debut = int(payload.get("debut"))
        fin = int(payload.get("fin"))
    except (TypeError, ValueError):
        raise ValueError("Invalid time slot.")
    if not (H_MIN <= debut < fin <= H_MAX):
        raise ValueError(f"Invalid time slot (between {H_MIN}h and {H_MAX}h, end after start).")

    try:
        effectif = int(float(payload.get("effectif", 0)))
    except (TypeError, ValueError):
        effectif = 0

    course = {
        "id": uuid.uuid4().hex[:10],
        "num_edt": num_edt, "filiere": filiere, "jour": jour,
        "debut": debut, "fin": fin, "duree": fin - debut,
        "module": module, "prof": prof, "effectif": max(effectif, 0),
        "salle": None, "_priority": _detect_priority(module),
    }
    cours.append(course)
    return course


def update_course(cours, course_id, payload):
    """Modifie un cours existant (identifie par son id) en place."""
    for c in cours:
        if c.get("id") != course_id:
            continue
        if "num_edt" in payload and str(payload["num_edt"]).strip():
            c["num_edt"] = str(payload["num_edt"]).strip()
        if "filiere" in payload and str(payload["filiere"]).strip():
            c["filiere"] = str(payload["filiere"]).strip()
        if "module" in payload and str(payload["module"]).strip():
            c["module"] = str(payload["module"]).strip()
            c["_priority"] = _detect_priority(c["module"])
        if "prof" in payload and str(payload["prof"]).strip():
            c["prof"] = str(payload["prof"]).strip()
        if "jour" in payload:
            jour = str(payload["jour"]).strip()
            if jour not in JOURS:
                raise ValueError("Invalid day.")
            c["jour"] = jour
        if "debut" in payload and "fin" in payload:
            try:
                d, f = int(payload["debut"]), int(payload["fin"])
            except (TypeError, ValueError):
                raise ValueError("Invalid time slot.")
            if not (H_MIN <= d < f <= H_MAX):
                raise ValueError(f"Invalid time slot (between {H_MIN}h and {H_MAX}h, end after start).")
            c["debut"], c["fin"], c["duree"] = d, f, f - d
        if "effectif" in payload:
            try:
                c["effectif"] = max(int(float(payload["effectif"])), 0)
            except (TypeError, ValueError):
                pass
        c["salle"] = None  # une modification invalide l'affectation precedente
        return c
    raise ValueError("Course not found.")


def remove_course(cours, course_id):
    """Supprime un cours (identifie par son id) de la liste, en place."""
    for i, c in enumerate(cours):
        if c.get("id") == course_id:
            return cours.pop(i)
    raise ValueError("Course not found.")


# ============================================
# MOTEUR (exécution en arrière-plan, pilotable)
# ============================================
class Moteur:
    def __init__(self, cours, pop, gen, pc, pm, cb_log=None, cb_progress=None, cb_done=None):
        self.cours = cours
        self.pop = pop
        self.gen = gen
        self.pc = pc
        self.pm = pm
        self.cb_log = cb_log or (lambda m: None)
        self.cb_progress = cb_progress or (lambda g, f, s: None)
        self.cb_done = cb_done or (lambda c, st, h: None)
        self._stop_event = threading.Event()
        # Wall-clock timing phase: records when the run starts/ends so that
        # execution time is measured and reported (not just logged as
        # "Not recorded" as it was before).
        self.start_time = None
        self.end_time = None
        self.elapsed_seconds = None

    def stop(self):
        self._stop_event.set()

    def run(self):
        self.start_time = time.time()
        self.cb_log(f"Starting: {len(self.cours)} courses, population={self.pop}, "
                     f"generations={self.gen}, crossover_rate={self.pc}, mutation_rate={self.pm}")

        course_objs = []
        for c in self.cours:
            co = core.Course(
                c["num_edt"], c["jour"], f'{c["debut"]}h-{c["fin"]}h',
                f'{c["module"]} - {c["prof"]}', c["effectif"])
            co.set_priority(c.get("_priority", 5))
            course_objs.append(co)

        rooms = core.load_or_init_rooms()
        params = {"population_size": self.pop, "number_generations": self.gen,
                  "mutation_rate": self.pm}
        ga = core.EvolutionaryAllocator(course_objs, rooms, params=params)

        # Respecte la probabilite de croisement demandee (pc) : sinon les
        # deux parents passent tels quels a la mutation.
        base_crossover = ga.crossover

        def _crossover(p1, p2):
            if random.random() > self.pc:
                return p1.copy(), p2.copy()
            return base_crossover(p1, p2)
        ga.crossover = _crossover

        state = {"best_seen": -1.0, "stag": 0}
        log_every = max(1, self.gen // 20)

        def _cb(gen_idx, best, avg):
            if best > state["best_seen"]:
                state["best_seen"] = best
                state["stag"] = 0
            else:
                state["stag"] += 1
            self.cb_progress(gen_idx + 1, best, state["stag"])
            if gen_idx == 0 or (gen_idx + 1) % log_every == 0:
                self.cb_log(f"Generation {gen_idx + 1}/{self.gen} — "
                             f"best={best:.1f}  avg={avg:.1f}  stagnation={state['stag']}")

        def _stop_flag():
            return self._stop_event.is_set()

        results, _, best_hist, avg_hist = ga.allocate(callback=_cb, stop_flag=_stop_flag)

        # Verification reelle des contraintes hard/soft sur le resultat
        # final de cette execution (compte les violations effectivement
        # presentes dans la solution retenue — pas une simple description).
        constraints_report = ga.verify_constraints(results)

        # results[i] correspond a self.cours[i] (meme ordre que course_objs)
        for i, r in enumerate(results):
            self.cours[i]["salle"] = r["room"].name

        self.end_time = time.time()
        self.elapsed_seconds = round(self.end_time - self.start_time, 3)

        n_over = sum(1 for r in results if r["course"].students > r["room"].capacity)
        rooms_used = sorted(set(r["room"].name for r in results))

        # ---- Statistical panel: dispersion of the fitness function ----
        # Variance / standard deviation of the fitness values inside the
        # final generation's population (diversity / convergence of the
        # fitness function at the end of the run), plus the same measures
        # computed over the best-fitness trajectory across all generations
        # (stability of the fitness function's progress during the run).
        final_pop_variance = round(ga.variance_fitness_history[-1], 3) if ga.variance_fitness_history else 0.0
        final_pop_stdev = round(ga.stdev_fitness_history[-1], 3) if ga.stdev_fitness_history else 0.0
        trajectory_variance = round(statistics.pvariance(best_hist), 3) if len(best_hist) > 1 else 0.0
        trajectory_stdev = round(statistics.pstdev(best_hist), 3) if len(best_hist) > 1 else 0.0

        stats = {
            "total_courses": len(self.cours),
            "allocated": len(results),
            "rooms_used": len(rooms_used),
            "rooms_total": len(rooms),
            "capacity_overbookings": n_over,
            "final_fitness": round(best_hist[-1], 1) if best_hist else 0,
            "execution_time_seconds": self.elapsed_seconds,
            "generations_run": len(best_hist),
            # Statistical panel (all English, as required by the UI)
            "fitness_variance": final_pop_variance,
            "fitness_stdev": final_pop_stdev,
            "fitness_trajectory_variance": trajectory_variance,
            "fitness_trajectory_stdev": trajectory_stdev,
            # Panneau contraintes : violations reellement comptees sur le
            # resultat de ce run (voir EvolutionaryAllocator.verify_constraints).
            "constraints": constraints_report,
        }
        if self._stop_event.is_set():
            self.cb_log(f"Stopped by user — {stats['allocated']} courses allocated "
                         f"(fitness {stats['final_fitness']}) in {self.elapsed_seconds:.1f}s.")
        else:
            self.cb_log(f"Done — final fitness {stats['final_fitness']}, "
                         f"{stats['allocated']} courses allocated, {n_over} capacity overbookings, "
                         f"completed in {self.elapsed_seconds:.1f}s "
                         f"({stats['generations_run']} generations run). "
                         f"Fitness variance: {stats['fitness_variance']} "
                         f"(std dev: {stats['fitness_stdev']}).")
        self.cb_done(self.cours, stats, best_hist)


# ============================================
# EXPORT EXCEL — reprend la charte graphique du site
# (bandeau bleu marine « ink », accents or « gold », occupation en vert)
# ============================================
_XL_INK = "142A42"
_XL_GOLD = "B4872F"
_XL_GOLD_SOFT = "F1E6CC"
_XL_FOREST_SOFT = "E4EFE6"
_XL_PANEL_LINE = "E4DCC8"
_XL_WHITE = "FFFFFF"

_thin = Side(style="thin", color=_XL_PANEL_LINE)
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _style_grid_sheet(ws, n_cols, header_row):
    """Style d'une grille (en-tete sombre, bordures, cellules centrees) a
    partir de la ligne `header_row` (1-indexee)."""
    for col in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 17
    for row in ws.iter_rows(min_row=header_row):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color=_XL_WHITE, name="Calibri", size=10.5)
        cell.fill = PatternFill(start_color=_XL_INK, end_color=_XL_INK, fill_type="solid")
    ws.row_dimensions[header_row].height = 22
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)


def _write_title_block(ws, n_cols, filiere_title, subtitle=""):
    """Ecrit le bandeau d'en-tete academique (universite / titre / sous-titre)
    en tete de feuille, avec le nom de la filiere bien mis en avant au
    centre, dans le meme esprit graphique que l'interface web."""
    last_col = get_column_letter(n_cols)

    ws.merge_cells(f"A1:{last_col}1")
    c1 = ws["A1"]
    c1.value = UNIVERSITE.upper()
    c1.font = Font(bold=True, color=_XL_GOLD_SOFT, size=10, name="Calibri")
    c1.fill = PatternFill(start_color=_XL_INK, end_color=_XL_INK, fill_type="solid")
    c1.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"A2:{last_col}2")
    c2 = ws["A2"]
    c2.value = filiere_title
    c2.font = Font(bold=True, color=_XL_WHITE, size=15, name="Calibri")
    c2.fill = PatternFill(start_color=_XL_INK, end_color=_XL_INK, fill_type="solid")
    c2.alignment = Alignment(horizontal="center", vertical="center")

    row3 = 3
    if subtitle:
        ws.merge_cells(f"A3:{last_col}3")
        c3 = ws["A3"]
        c3.value = subtitle
        c3.font = Font(italic=True, color=_XL_INK, size=10.5, name="Calibri")
        c3.fill = PatternFill(start_color=_XL_GOLD_SOFT, end_color=_XL_GOLD_SOFT, fill_type="solid")
        c3.alignment = Alignment(horizontal="center", vertical="center")
        row3 = 4

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 26
    if subtitle:
        ws.row_dimensions[3].height = 18
    return row3  # premiere ligne libre apres le bandeau


def build_workbooks(cours):
    """Retourne (classeur_EDT, classeur_salles, horodatage). Chaque feuille
    reprend la charte graphique de l'interface web (bandeau bleu marine /
    or) avec le nom de la filiere bien visible en tete de grille."""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    n_cols = len(HEURES) + 1

    # ---------- Classeur 1 : Emploi du temps par filiere ----------
    wb1 = Workbook()
    wb1.remove(wb1.active)
    edts = sorted(set(c["num_edt"] for c in cours if c.get("salle")))
    for edt in edts:
        gc = [c for c in cours if c["num_edt"] == edt and c.get("salle")]
        filiere = gc[0]["filiere"] if gc else edt
        ws = wb1.create_sheet(title=str(edt)[:31])

        header_row = _write_title_block(
            ws, n_cols, filiere_title=filiere,
            subtitle=f"Schedule — Group / EDT: {edt}")

        ws.cell(row=header_row, column=1, value="Day")
        for i, h in enumerate(HEURES, start=2):
            ws.cell(row=header_row, column=i, value=f"{h}h-{h + 1}h")

        for r, jour in enumerate(JOURS, start=header_row + 1):
            ws.cell(row=r, column=1, value=jour).font = Font(bold=True, size=10.5)
            ws.cell(row=r, column=1).fill = PatternFill(
                start_color=_XL_GOLD_SOFT, end_color=_XL_GOLD_SOFT, fill_type="solid")
            for i, h in enumerate(HEURES, start=2):
                found = next((c for c in gc if c["jour"] == jour and c["debut"] <= h < c["fin"]), None)
                cell = ws.cell(row=r, column=i)
                if found:
                    cell.value = f'{found["module"]}\n{found["prof"]} · {found["salle"]}'
                    cell.fill = PatternFill(start_color=_XL_FOREST_SOFT, end_color=_XL_FOREST_SOFT, fill_type="solid")
        _style_grid_sheet(ws, n_cols, header_row)
        ws.row_dimensions[header_row].height = 22
        for r in range(header_row + 1, header_row + 1 + len(JOURS)):
            ws.row_dimensions[r].height = 34
    if not wb1.sheetnames:
        ws = wb1.create_sheet(title="No results")
        ws["A1"] = "No courses allocated — run the allocation first."

    # ---------- Classeur 2 : Occupation des salles ----------
    wb2 = Workbook()
    wb2.remove(wb2.active)
    amphi, moyenne, petite, salles_all, capacites = get_salles()
    for jour in ["Toutes"] + JOURS:
        sn = "All_days" if jour == "Toutes" else jour
        ws = wb2.create_sheet(title=sn[:31])
        sous_titre = "Room occupancy — all days" if jour == "Toutes" \
            else f"Room occupancy — {jour}"
        header_row = _write_title_block(ws, n_cols, filiere_title=UNIVERSITE, subtitle=sous_titre)

        ws.cell(row=header_row, column=1, value="Room")
        for i, h in enumerate(HEURES, start=2):
            ws.cell(row=header_row, column=i, value=f"{h}h-{h + 1}h")

        r = header_row
        for salle in sorted(salles_all):
            r += 1
            label = f"{salle} ({capacites.get(salle, '?')} pl.)"
            ws.cell(row=r, column=1, value=label).font = Font(bold=True, size=10.5)
            ws.cell(row=r, column=1).fill = PatternFill(
                start_color=_XL_GOLD_SOFT, end_color=_XL_GOLD_SOFT, fill_type="solid")
            for i, h in enumerate(HEURES, start=2):
                found = next((c for c in cours if c.get("salle") == salle
                              and (jour == "Toutes" or c["jour"] == jour)
                              and c["debut"] <= h < c["fin"]), None)
                cell = ws.cell(row=r, column=i)
                if found:
                    cell.value = f'{found["module"]}\n{found["num_edt"]} · {found["prof"]}'
                    cell.fill = PatternFill(start_color=_XL_FOREST_SOFT, end_color=_XL_FOREST_SOFT, fill_type="solid")
        _style_grid_sheet(ws, n_cols, header_row)
        for row in range(header_row + 1, r + 1):
            ws.row_dimensions[row].height = 34

    return wb1, wb2, ts
